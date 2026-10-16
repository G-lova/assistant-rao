"""Тесты форм объекта 7 (отчётные документы по контракту/договору): разбор листа Excel, выбор формы, черновик."""
import asyncio
import json
import os
import unittest

import openpyxl

from knowledge_store import forms, summary
from tests import pg_schema
from tests.pg_psql import PsqlConn
from tests.test_summary import S

PG = os.getenv("PE_TEST_PG")


def make_sheet(rows):
    """Строит лист в раскладке Excel: колонки F–I, «Объект 7 (14)» в I1.

    Args:
        rows: Кортежи ``(раздел, подраздел, название, ключ)``; пустое значение — пустая ячейка.

    Returns:
        Лист ``openpyxl``.
    """
    ws = openpyxl.Workbook().active
    ws["I1"] = "Объект 7 (14)"
    ws["F2"] = "Поставка товаров Новые"
    for n, (section, sub, label, key) in enumerate(rows, start=3):
        for col, value in zip("FGHI", (section, sub, label, key)):
            if value:
                ws[f"{col}{n}"] = value
    return ws


ROWS = [
    ("Раздел 1. Документы и материалы", None, "Укажите перечень документов и материалов", "field1_1"),
    (None, None, "Наименование товаров", "field1_3[0]['name']"),
    (None, None, "Требования к составу отчетной документации", "field1_3[0]['requirements']"),
    ("Раздел 3. Оценка", "Подраздел 3.3. Сроки", "Не представляется возможным определить сроки", "field3_0_3 (bool)"),
    (None, None, "Количество календарных дней просрочки", "field3_0_5"),
    ("Раздел 5. Заключение", "Подраздел 5.1. Расчет неустойки", "Ключевая ставка", "field6_1_1"),
    (None, None, "Параметры состава", "field1_3[0]['requirements']"),             # повтор ключа
    ("6. Документация по приемке", None, "Представлена документация по приемке?", "field7_0"),
    (None, None, "Какой-то заголовок без ключа", None),
]


class ReportFormsParseTests(unittest.TestCase):
    """Разбор листа объекта 7."""

    def setUp(self):
        """Разбирает тестовый лист."""
        self.reg = forms.parse_report_sheet(make_sheet(ROWS), "goods")
        self.fields = {f.field_key: f for f in self.reg["44fz_goods_obj7"]}

    def test_both_laws_get_same_fields(self):
        """Лист «223-фз и 44-фз» даёт формы обоих законов с одинаковыми полями."""
        self.assertEqual(set(self.reg), {"44fz_goods_obj7", "223fz_goods_obj7"})
        self.assertEqual([f.field_key for f in self.reg["44fz_goods_obj7"]], [f.field_key for f in self.reg["223fz_goods_obj7"]])

    def test_keys_and_duplicates(self):
        """Ключи массивных полей сохраняются, повтор ключа остаётся за первым критерием, строки без ключа пропускаются."""
        self.assertIn("field1_3[0]['name']", self.fields)
        self.assertEqual(self.fields["field1_3[0]['requirements']"].label, "Требования к составу отчетной документации")
        self.assertEqual(sum(1 for f in self.reg["44fz_goods_obj7"] if f.field_key == "field1_3[0]['requirements']"), 1)
        self.assertEqual(list(forms.report_duplicates(make_sheet(ROWS))), ["field1_3[0]['requirements']"])

    def test_kinds_and_levels(self):
        """Пометка (bool) и вопросы — bool, расчёты — number, раздел 1 проверяет код, остальное эксперт."""
        self.assertEqual(self.fields["field3_0_3"].value_kind, forms.KIND_BOOL)
        self.assertEqual(self.fields["field7_0"].value_kind, forms.KIND_BOOL)
        self.assertEqual(self.fields["field6_1_1"].value_kind, forms.KIND_NUMBER)
        self.assertEqual(self.fields["field3_0_5"].value_kind, forms.KIND_NUMBER)
        self.assertEqual(self.fields["field1_1"].value_kind, forms.KIND_TEXT)
        self.assertEqual((self.fields["field1_1"].check_level, self.fields["field1_3[0]['name']"].check_level), (1, 1))
        self.assertEqual(self.fields["field3_0_5"].check_level, 3)

    def test_section_path_and_service_field(self):
        """Раздел/подраздел сохраняются, служебное поле комментария достроено."""
        self.assertEqual(self.fields["field3_0_3"].section, "Раздел 3. Оценка / Подраздел 3.3. Сроки")
        self.assertTrue(self.fields["rao_comment"].derived)

    def test_sheet_without_block(self):
        """Лист без «Объект 7» формы не даёт."""
        self.assertEqual(forms.parse_report_sheet(openpyxl.Workbook().active, "nir"), {})


class ReportTitleTests(unittest.TestCase):
    """Эвристика названий распознаёт отчётные документы."""

    def test_act_and_report_titles(self):
        """Акт, отчёт, счёт-фактура и накладная берутся из первой строки."""
        from knowledge_store import doc_titles
        for head in ("Акт сдачи-приемки выполненных работ № 5 от 01.02.2026\nг. Москва",
                     "Отчет о выполнении НИР за 2025 год\nРуководитель",
                     "Счет-фактура № 12 от 03.03.2026\nПродавец",
                     "Товарная накладная № 7 от 03.03.2026"):
            self.assertEqual(doc_titles.heuristic_title(head), head.splitlines()[0].strip(" .;:"))
        self.assertIsNone(doc_titles.heuristic_title("Москва, март 2026\nактуальность работы"))


class ReportFormCodeTests(unittest.TestCase):
    """Выбор формы объекта 7 по checkType2."""

    def test_kinds_by_check_type(self):
        """13 — НИР, 14 — поставка, 15 — работы/услуги; закон входит в код."""
        self.assertEqual(forms.get_form_code("44-ФЗ", 13, 7), "44fz_nir_obj7")
        self.assertEqual(forms.get_form_code("223-ФЗ", 14, 7), "223fz_goods_obj7")
        self.assertEqual(forms.get_form_code("44-ФЗ", 15, 7), "44fz_works_obj7")

    def test_other_combinations_unsupported(self):
        """Способы закупки с объектом 7 и отчётные виды с объектом 6 формы не имеют."""
        self.assertIsNone(forms.get_form_code("44-ФЗ", 1, 7))
        self.assertIsNone(forms.get_form_code("44-ФЗ", 13, 6))

    def test_available_filter(self):
        """Не загруженная в справочник форма не поддерживается."""
        self.assertIsNone(forms.get_form_code("44-ФЗ", 13, 7, available=["44fz_goods_obj7"]))
        self.assertTrue(forms.is_report_form("223fz_works_obj7"))
        self.assertFalse(forms.is_report_form("44fz_competition_obj6"))


@unittest.skipUnless(PG, "нужен PE_TEST_PG=host:port")
class ReportSummaryDbTests(unittest.TestCase):
    """Черновик заключения по отчётным документам на настоящем Postgres."""

    @classmethod
    def setUpClass(cls):
        """Создаёт схему pe_*."""
        host, port = PG.split(":")
        cls.conn = PsqlConn(host, int(port))
        pg_schema.build(cls.conn)

    def test_draft_has_only_document_list(self):
        """Факты не нужны; заполняется перечень документов, остальное пусто, статус needs_review."""
        c, form = self.conn, "44fz_works_obj7"
        c._run(f"DELETE FROM pe_form_fields WHERE form_code = '{form}'")
        for n, (key, kind) in enumerate([("field1_1", "text"), ("field3_0_3", "bool"), ("field6_1_1", "number"),
                                          ("rao_comment", "text")], start=1):
            c._run(f"INSERT INTO pe_form_fields (form_code, field_key, ordinal, label, value_kind) "
                   f"VALUES ('{form}', '{key}', {n}, '{key}', '{kind}')")
        c._run("INSERT INTO pe_procurements (expertise_id, law, method, object_code, check_type2) "
               "VALUES (8301, '44-ФЗ', 'Работы', 7, 15)")
        c._run("INSERT INTO pe_documents (expertise_id, doc_code, filename, sha256, text_full) "
               "VALUES (8301, 'docOther', 'akt.docx', 'h8301', 'Акт сдачи-приемки выполненных работ № 5 от 01.02.2026\nг. Москва')")
        res = asyncio.run(summary.generate_summary(c, S(8301), None))
        self.assertEqual(res["form_code"], form)
        self.assertEqual(res["status"], "needs_review")
        self.assertEqual(set(res["data"]), {"field1_1", "field3_0_3", "field6_1_1", "rao_comment"})
        self.assertIsNone(res["data"]["field3_0_3"])
        self.assertIn("акт", res["data"]["field1_1"].lower())
        stored = json.loads(c._run("SELECT data FROM pe_summary_opinions WHERE expertise_id = 8301"))
        self.assertEqual(stored, res["data"])
        c._run(f"DELETE FROM pe_form_fields WHERE form_code = '{form}'")


if __name__ == "__main__":
    unittest.main()
