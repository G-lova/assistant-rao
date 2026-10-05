"""Тесты этапа 3: справочник форм из Excel, выбор формы, запись в ``pe_form_fields``.

Excel собирается в памяти по раскладке листа «Все Закупки 44-фз» (блок = 6 колонок, ключи в
колонке «Объект N»), поэтому тесты не зависят от реального файла. Интеграционные тесты идут на
настоящем Postgres, если задана переменная ``PE_TEST_PG=host:port``.
"""
import asyncio
import os
import re
import unittest
from pathlib import Path

import openpyxl

from knowledge_store import forms, repository as repo
from tests.pg_psql import PsqlConn

PG = os.getenv("PE_TEST_PG")
MIGRATIONS = Path(__file__).resolve().parent.parent / "migrations"


def make_sheet():
    """Строит лист в раскладке «Все Закупки 44-фз» с двумя формами (конкурс и аукцион)."""
    ws = openpyxl.Workbook().active
    # шапка: название формы в колонках A и G, «Объект 6» в E и K
    ws.cell(1, 1, "Конкурс новый 44-ФЗ"); ws.cell(1, 5, "Объект 6")
    ws.cell(1, 7, "Аукцион Новый 44-ФЗ"); ws.cell(1, 11, "Объект 6")
    rows = [
        (2, "Блок Экспертное заключение", None, None, None, None),
        (3, None, None, "Шифр проекта", "code", None),
        (4, "Блок I. Общие сведения", None, None, None, None),
        (5, None, None, "Начальная цена, руб", "field1_1", None),
        (6, "Блок II", None, None, None, None),
        (7, None, "Раздел 1. Экспертная оценка", None, None, None),
        (8, None, None, "1.1. Наличие информации о наименовании Заказчика", "field2_1_1", None),
        (9, None, None, "Комментарий к разделу", "field2_1_text", None),
        (10, None, "Раздел 2. Соответствие", None, None, None),
        (11, None, None, "Подраздел 2.1. Описание", None, None),
        (12, None, None, "2.1. Соответствие описания объекта", "field2_2_1", None),
        (13, None, None, "2.1.1. Соответствие наименований", "field2_2_1_1", None),
        (14, None, None, "Подраздел 2.4. Порядок оценки", None, None),
        (15, None, None, "2.4.1. Соответствие порядка оценки", "field2_2_4_1", None),
        (16, None, None, "Выберите метод обоснования", "field2_2_2_0", None),
        (17, "Блок III. Вывод", None, None, None, None),
        (18, None, None, "Вывод", "field3", None),
        (19, "Блок IV. Заключение", None, None, None, None),
        (20, None, None, "Заключение", "field4", None),
    ]
    for r, a, b, label, key, _ in rows:
        ws.cell(r, 1, a); ws.cell(r, 2, b)
        ws.cell(r, 4, label); ws.cell(r, 5, key)
    # второй блок (аукцион): те же строки, но другая нумерация и сдвиг колонок на 6
    ws.cell(3, 10, "Шифр проекта"); ws.cell(3, 11, "code")
    ws.cell(8, 10, "1.1. Наличие информации о наименовании Заказчика"); ws.cell(8, 11, "field2_1_1")
    ws.cell(12, 10, "2.1. Соответствие чего-то"); ws.cell(12, 11, "field2_2_0")
    return ws


class ParseTests(unittest.TestCase):
    """Разбор листа Excel в справочник полей."""

    @classmethod
    def setUpClass(cls):
        """Разбирает синтетический лист один раз."""
        cls.reg = forms.parse_forms_sheet(make_sheet())

    def keys(self, code):
        """Ключи полей формы в порядке ordinal."""
        return [f.field_key for f in self.reg[code]]

    def test_forms_found_and_codes(self):
        """Формы распознаются по названию, код = закон_способ_объект."""
        self.assertEqual(set(self.reg), {"44fz_competition_obj6", "44fz_auction_obj6"})

    def test_keys_come_from_excel_only(self):
        """У аукциона нет ключей конкурса: нумерация берётся из своего блока, а не хардкодится."""
        auction = set(self.keys("44fz_auction_obj6"))
        self.assertIn("field2_2_0", auction)
        self.assertNotIn("field2_2_1_1", auction)
        self.assertNotIn("field2_2_4_1", auction)

    def test_order_and_ordinals(self):
        """Порядок полей — как в Excel, ordinal сквозной и без дублей."""
        comp = self.reg["44fz_competition_obj6"]
        self.assertEqual([f.ordinal for f in comp], list(range(1, len(comp) + 1)))
        self.assertLess(self.keys("44fz_competition_obj6").index("field2_1_1"),
                        self.keys("44fz_competition_obj6").index("field4"))
        self.assertEqual(len(self.keys("44fz_competition_obj6")), len(set(self.keys("44fz_competition_obj6"))))

    def test_kinds(self):
        """Типы значений: наличие, соответствие, итог подраздела, выбор, текст, мета."""
        kind = {f.field_key: f.value_kind for f in self.reg["44fz_competition_obj6"]}
        self.assertEqual(kind["code"], forms.KIND_META)
        self.assertEqual(kind["field1_1"], forms.KIND_NUMBER)
        self.assertEqual(kind["field2_1_1"], forms.KIND_PRESENCE)
        self.assertEqual(kind["field2_1_text"], forms.KIND_TEXT)
        self.assertEqual(kind["field2_2_1"], forms.KIND_SECTION)
        self.assertEqual(kind["field2_2_1_1"], forms.KIND_COMPLIANCE)
        self.assertEqual(kind["field2_2_2_0"], forms.KIND_CHOICE)
        self.assertEqual(kind["field3"], forms.KIND_TEXT)

    def test_derived_keys(self):
        """Достраиваются <ключ>_text, родитель подраздела без строки в Excel и служебные поля."""
        comp = {f.field_key: f for f in self.reg["44fz_competition_obj6"]}
        self.assertTrue(comp["field2_2_1_1_text"].derived)
        self.assertTrue(comp["field2_2_4"].derived)            # подраздел 2.4 без своей строки
        self.assertTrue(comp["field2_2_4_text"].derived)
        self.assertTrue(comp["rao_comment"].derived)
        self.assertFalse(comp["field2_1_text"].derived)        # есть в Excel — не производный
        self.assertNotIn("field2_1_1_text", comp)              # для раздела 1 комментариев по критериям нет

    def test_sections_are_recorded(self):
        """Поле знает свой блок/раздел/подраздел."""
        comp = {f.field_key: f for f in self.reg["44fz_competition_obj6"]}
        self.assertIn("Блок II", comp["field2_2_1_1"].section)
        self.assertIn("Раздел 2", comp["field2_2_1_1"].section)


class FormCodeTests(unittest.TestCase):
    """Выбор формы по полям экспертизы (шаг 3.4)."""

    def test_competition_44fz_obj6(self):
        """checkType2 1..3 → конкурс."""
        for ct in (1, 2, 3):
            self.assertEqual(forms.get_form_code("44-ФЗ", ct, 6), "44fz_competition_obj6")

    def test_other_methods(self):
        """Аукцион, котировки, единственный поставщик."""
        self.assertEqual(forms.get_form_code("44-ФЗ", 5, 6), "44fz_auction_obj6")
        self.assertEqual(forms.get_form_code("44-ФЗ", 7, 6), "44fz_quotation_obj6")
        self.assertEqual(forms.get_form_code("44-ФЗ", 11, 6), "44fz_single_supplier_obj6")

    def test_unsupported_returns_none(self):
        """Неизвестные значения и пустые данные не ломают, а дают None."""
        self.assertIsNone(forms.get_form_code("", 1, 6))
        self.assertIsNone(forms.get_form_code("44-ФЗ", 9, 6))
        self.assertIsNone(forms.get_form_code("44-ФЗ", None, 6))
        self.assertIsNone(forms.get_form_code("44-ФЗ", 1, float("nan")))

    def test_available_filter(self):
        """Форма вне справочника (например, объект 7) пока не поддерживается."""
        self.assertIsNone(forms.get_form_code("44-ФЗ", 1, 7, available={"44fz_competition_obj6"}))
        self.assertEqual(forms.get_form_code("44-ФЗ", 1, 6, available={"44fz_competition_obj6"}), "44fz_competition_obj6")


@unittest.skipUnless(PG, "нужен PE_TEST_PG=host:port")
class FormFieldsDbTests(unittest.TestCase):
    """Запись справочника в настоящий Postgres."""

    @classmethod
    def setUpClass(cls):
        """Создаёт таблицу pe_form_fields из миграций 001 и 002."""
        host, port = PG.split(":")
        cls.conn = PsqlConn(host, int(port))
        sql = (MIGRATIONS / "001_knowledge_store.sql").read_text(encoding="utf-8")
        sql = re.sub(r"CREATE EXTENSION[^\n]*\n", "", sql)
        sql = re.sub(r"[^\n]*USING hnsw[^\n]*\n", "", sql).replace("vector(1024)", "real[]")
        cls.conn._run("DROP TABLE IF EXISTS pe_summary_opinions, pe_facts, pe_chunks, pe_documents, pe_form_fields, pe_procurements CASCADE;")
        cls.conn._run(sql)
        cls.conn._run((MIGRATIONS / "002_form_fields_derived.sql").read_text(encoding="utf-8"))
        cls.reg = forms.parse_forms_sheet(make_sheet())

    def test_replace_is_idempotent(self):
        """Повторный импорт не плодит дубли и сохраняет флаг derived."""
        code = "44fz_competition_obj6"
        n1 = asyncio.run(repo.replace_form_fields(self.conn, code, self.reg[code]))
        n2 = asyncio.run(repo.replace_form_fields(self.conn, code, self.reg[code]))
        self.assertEqual(n1, n2)
        self.assertEqual(self.conn._run(f"SELECT count(*) FROM pe_form_fields WHERE form_code='{code}'"), str(n1))
        self.assertEqual(self.conn._run(
            f"SELECT derived FROM pe_form_fields WHERE form_code='{code}' AND field_key='rao_comment'"), "t")

    def test_read_back_and_codes(self):
        """Поля читаются в порядке ordinal, коды форм перечисляются."""
        for code, fields in self.reg.items():
            asyncio.run(repo.replace_form_fields(self.conn, code, fields))
        codes = asyncio.run(repo.list_form_codes(self.conn))
        self.assertEqual(set(codes), set(self.reg))
        rows = asyncio.run(repo.get_form_fields(self.conn, "44fz_competition_obj6"))
        self.assertEqual([r["field_key"] for r in rows], [f.field_key for f in self.reg["44fz_competition_obj6"]])
        self.assertEqual(rows[0]["field_key"], "code")


if __name__ == "__main__":
    unittest.main()
