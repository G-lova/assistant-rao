"""Чтение .xlsx с ограничениями (``configs.file_reader.xlsx_text_limited``): большая смета не подвешивает прогон."""
import os
import tempfile
import time
import unittest

from tests.stubs import stubbed_import


class ExcelLimitedTest(unittest.TestCase):
    """Потоковое чтение: строки, столбцы, листы и пустой «хвост» листа ограничены."""

    def setUp(self):
        """Функция модуля без тяжёлых зависимостей (заглушки действуют только на время импорта)."""
        with stubbed_import("configs.file_reader") as module:
            self.read = module.xlsx_text_limited

    def book(self, rows: int, empty_tail: bool = False) -> str:
        """Временная книга: лист «Смета» с ``rows`` строками по 30 столбцов, второй лист — короткий."""
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = "Смета"
        for r in range(1, rows + 1):
            ws.append([f"Позиция {r}", r] + [f"т{c}" for c in range(28)])
        if empty_tail:
            for r in range(rows + 1, rows + 30001):
                ws.cell(row=r, column=1).font = Font(bold=True)   # оформленные пустые строки
        wb.create_sheet("Итоги").append(["Итого", 100])
        f = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        f.close()
        wb.save(f.name)
        self.addCleanup(os.unlink, f.name)
        return f.name

    def test_limits(self):
        """Не больше max_rows строк и max_cols столбцов, лист помечен обрезанным, следующий лист прочитан."""
        t0 = time.time()
        text = self.read(self.book(5000), max_rows=100, max_cols=5)
        self.assertLess(time.time() - t0, 10)
        self.assertIn("Лист: Смета", text)
        self.assertIn("Позиция 100 |", text)
        self.assertNotIn("Позиция 101 |", text)
        self.assertIn("[Лист обрезан: прочитано 100 строк]", text)
        self.assertNotIn("т3", text)          # 6-й столбец и дальше не читаются
        self.assertIn("Лист: Итоги\nИтого | 100", text)

    def test_empty_tail_and_chars(self):
        """Оформленные пустые строки не читаются до конца листа; общий объём текста ограничен."""
        text = self.read(self.book(10, empty_tail=True), max_rows=50)
        self.assertIn("Позиция 10 |", text)
        self.assertIn("Лист: Итоги", text)
        self.assertLessEqual(len(self.read(self.book(2000), max_chars=1000)), 1000)


if __name__ == "__main__":
    unittest.main()
